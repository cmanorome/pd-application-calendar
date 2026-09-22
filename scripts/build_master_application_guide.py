#!/usr/bin/env python3
"""Join catalog, calendar intervals, usage-guide notes, and rates into one workbook.

Outputs:
  data/pd_master_products.csv
  data/pd_master_program_notes.csv
  data/pd_master_application_guide.xlsx

Join key is sku (uppercase). Load CSV with:
  csv.DictReader(open(path, encoding="utf-8-sig"))
Treat 0/1 flag columns as booleans.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:  # pragma: no cover
    Workbook = None  # type: ignore

from pd_engine.catalog import Catalog

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"

FLAG_COLS = (
    "mix_together",
    "independent",
    "water_in",
    "soil_drench",
    "foliar",
    "fertigation",
    "hydroponic",
    "hose_on",
    "mix_with_iron",
    "calendar_one_off",
    "garden_natives_half_strength",
    "in_catalog",
    "in_calendar",
    "in_usage_guide",
    "in_rates_calculator",
)

PRODUCT_COLS = [
    "sku",
    "catalog_name",
    "guide_name",
    "role_type",
    "form",
    "category",
    "analysis",
    "npk",
    "product_url",
    "pack_sizes",
    "calendar_how_often",
    "calendar_cadence_days",
    "calendar_window",
    "calendar_one_off",
    "calendar_max_per_year",
    "calendar_tank_group",
    "calendar_method",
    "calendar_notes",
    "mix_together",
    "independent",
    "water_in",
    "soil_drench",
    "foliar",
    "fertigation",
    "hydroponic",
    "hose_on",
    "mix_with_iron",
    "usage_frequency_text",
    "usage_rate_text",
    "usage_min_dilution",
    "usage_application_notes",
    "calculator_frequency",
    "calculator_dilution",
    "calculator_rate_100m2",
    "calculator_notes",
    "garden_natives_half_strength",
    "combined_notes",
    "in_catalog",
    "in_calendar",
    "in_usage_guide",
    "in_rates_calculator",
]

NATIVES_HALF_SKUS = frozenset({"SWS", "A8M"})

PROGRAM_NOTES = [
    {
        "note_id": "winter_nights_below_10c",
        "applies_when": "Garden or lawn plans with year-round sprays (cadence 42 days or less)",
        "sku": "",
        "text": "Keep applying through winter, just less often. Winter means nights below 10°C — usually June to August in southern Australia, and shorter or sometimes missing in the north.",
    },
    {
        "note_id": "garden_natives_half_strength",
        "applies_when": "Garden selected",
        "sku": "SWS, A8M",
        "text": "For natives and other sensitive plants, use Seaweed Secrets and Activ8Mate at half strength.",
    },
    {
        "note_id": "humate_top_of_soil",
        "applies_when": "Humate granules (513) on the plan",
        "sku": "513",
        "text": "Spread Humate granules on the top layer of soil, under mulch, or dug in. They do not need to be watered in.",
    },
    {
        "note_id": "ffr_alternate_activ8",
        "applies_when": "Flowers, Fruits & Roots liquid plus Activ8Mate or Activ8EXTRA",
        "sku": "721, A8M, A8X",
        "text": "Alternate between Flowers, Fruits & Roots and Activ8 each feed. FFR is for flowering/fruiting; Activ8 is the regular feed. Do not mix those two in the same sprayer.",
    },
    {
        "note_id": "mixable_same_day",
        "applies_when": "Two or more mix_together concentrates",
        "sku": "",
        "text": "Mix mixable concentrates in one sprayer on the same day. Do not add iron to that tank. Jar test if it is a new combination.",
    },
    {
        "note_id": "iron_separate_day",
        "applies_when": "Liquid iron plus seaweed, humic, wetter, or Activ8",
        "sku": "LIR, LEN, MG, 547, 846",
        "text": "Liquid iron is on a different day from seaweed, humic, wetter, and Activ8. Do not mix them in the same sprayer.",
    },
    {
        "note_id": "lime_wait_before_iron",
        "applies_when": "Lime or dolomite and an iron product",
        "sku": "LIMEGr, DOL",
        "text": "Wait 42 days after lime or dolomite before applying iron.",
    },
    {
        "note_id": "iron_gap_days",
        "applies_when": "Iron with seaweed / humic / wetter / Activ8",
        "sku": "",
        "text": "Keep liquid iron at least 3 days away from seaweed, humic, wetter, and Activ8.",
    },
    {
        "note_id": "volcamin_no_badges",
        "applies_when": "VolcaMin (557)",
        "sku": "557",
        "text": "VolcaMin has no mix, apply-independently, or water-in badges. Incorporate into soil, potting mix, or the root zone.",
    },
    {
        "note_id": "stimulizer_iron_exception",
        "applies_when": "Tank-mixing",
        "sku": "STM",
        "text": "Stimulizer is the only product the usage guide says can be tank-mixed with iron.",
    },
    {
        "note_id": "default_home_rate",
        "applies_when": "No product-specific rate",
        "sku": "",
        "text": "Unless otherwise specified, apply 10 mL per 10 m² (100 mL per 100 m²) of lawn or garden.",
    },
    {
        "note_id": "source",
        "applies_when": "Always",
        "sku": "",
        "text": "Joined from Product Usage Guide (Aug 2025), PD application rates sheet, schedule_intervals.csv, and the product catalog. Follow the label.",
    },
]


def norm_sku(value: str) -> str:
    return (value or "").strip().upper()


def clean(value: object) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"[ \t]+", " ", text.replace("\xa0", " ")).strip()


def unique_headers(raw: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out: list[str] = []
    for item in raw:
        base = clean(item) or "col"
        n = seen.get(base, 0) + 1
        seen[base] = n
        out.append(base if n == 1 else f"{base} {n}")
    return out


def read_dicts(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        return [{k: clean(v) for k, v in row.items()} for row in csv.DictReader(f)]


def read_rates_calculator(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        raw_headers = next(reader, [])
        headers = unique_headers(raw_headers)
        rows: list[dict[str, str]] = []
        for raw in reader:
            if not raw or not any(c.strip() for c in raw):
                continue
            rec = {headers[i]: clean(raw[i]) if i < len(raw) else "" for i in range(len(headers))}
            sku = norm_sku(rec.get("SKU") or rec.get("sku") or "")
            if not sku:
                continue
            rec["sku"] = sku
            packs = [rec.get(h, "") for h in headers if h.lower().startswith("unit size") and rec.get(h)]
            rec["pack_sizes"] = ", ".join(packs)
            rec["product_url"] = rec.get("Buy now") or rec.get("Buy now 2") or ""
            rec["calculator_notes"] = rec.get("Notes") or ""
            rec["calculator_frequency"] = rec.get("Frequency") or ""
            rec["calculator_dilution"] = rec.get("Dilution Rate") or rec.get("Dilution Guide") or ""
            rec["rate_100m2"] = _calc_rate_100m2(rec)
            rows.append(rec)
    return rows


def _first(*values: str) -> str:
    for value in values:
        if clean(value):
            return clean(value)
    return ""


def _calc_rate_100m2(rec: dict[str, str]) -> str:
    ml = _range_label(
        rec.get("Rate Min mL / 100 m²"),
        rec.get("Rate Max mL / 100 m²"),
        "mL",
    )
    if ml:
        return f"{ml} / 100 m²"
    kg = _range_label(
        rec.get("Rate  Min kg / 100 m²"),
        rec.get("Rate  Max kg / 100 m²"),
        "kg",
    )
    if kg:
        return f"{kg} / 100 m²"
    g100 = clean(rec.get("Rate g / 100 m²"))
    if g100:
        return f"{g100} g / 100 m²"
    gm2 = _range_label(rec.get("Rate Min g / m²"), rec.get("Rate Max g / m²"), "g")
    if gm2:
        return f"{gm2} / m²"
    return ""


def _range_label(lo: str | None, hi: str | None, unit: str) -> str:
    a = clean(lo).replace(unit, "").strip()
    b = clean(hi).replace(unit, "").strip()
    if a and b and a != b:
        return f"{a}–{b} {unit}"
    if b:
        return f"{b} {unit}"
    if a:
        return f"{a} {unit}"
    return ""


def index_by_sku(rows: list[dict[str, str]], key: str = "sku") -> dict[str, dict[str, str]]:
    out: dict[str, dict[str, str]] = {}
    for row in rows:
        sku = norm_sku(row.get(key) or "")
        if sku and sku not in out:
            out[sku] = row
    return out


def group_by_sku(rows: list[dict[str, str]], key: str = "sku") -> dict[str, list[dict[str, str]]]:
    out: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        sku = norm_sku(row.get(key) or "")
        if not sku:
            continue
        out.setdefault(sku, []).append(row)
    return out


def pick_flag(interval: dict[str, str], usage: dict[str, str], key: str) -> str:
    if key in interval and interval.get(key) != "":
        return interval[key]
    if key in usage and usage.get(key) != "":
        return usage[key]
    return "0"


def first_nonempty(rows: list[dict[str, str]], *keys: str) -> str:
    for row in rows:
        for key in keys:
            val = clean(row.get(key))
            if val:
                return val
    return ""
    for row in rows:
        for key in keys:
            val = clean(row.get(key))
            if val:
                return val
    return ""


def join_unique(values: list[str]) -> str:
    seen: set[str] = set()
    out: list[str] = []
    for raw in values:
        text = clean(raw)
        if not text:
            continue
        key = re.sub(r"\s+", " ", text.lower())
        if key in seen:
            continue
        seen.add(key)
        out.append(text)
    return " | ".join(out)


def build_products() -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]]]:
    catalog = Catalog.from_csv(DATA / "plant_doctor_recommendation_engine_template.csv")
    catalog_by = {norm_sku(p.id): p for p in catalog.products if p.role_type.value != "bundle"}
    intervals = index_by_sku(read_dicts(DATA / "schedule_intervals.csv"))
    guide = index_by_sku(read_dicts(DATA / "product_usage_guide.csv"))
    guide_rates = read_dicts(DATA / "product_usage_guide_rates.csv")
    calc_rows = read_rates_calculator(DATA / "PD-application-rates-sheet.csv")
    calc_by = group_by_sku(calc_rows)

    skus = sorted(
        set(catalog_by) | set(intervals) | set(guide) | set(calc_by),
        key=lambda s: (s.isalpha() is False, s),
    )
    products: list[dict[str, str]] = []
    for sku in skus:
        product = catalog_by.get(sku)
        interval = intervals.get(sku, {})
        usage = guide.get(sku, {})
        calc = calc_by.get(sku, [])
        natives = "1" if sku in NATIVES_HALF_SKUS else "0"
        calendar_notes = interval.get("notes") or ""
        usage_notes = usage.get("application_notes") or ""
        calc_notes = join_unique([r.get("calculator_notes") or "" for r in calc])
        extra = []
        if natives == "1":
            extra.append(
                "For natives and other sensitive plants, use Seaweed Secrets and Activ8Mate at half strength."
            )
        combined = join_unique([calendar_notes, usage_notes, calc_notes, *extra])
        products.append(
            {
                "sku": product.id if product else (interval.get("sku") or usage.get("sku") or sku),
                "catalog_name": product.name if product else "",
                "guide_name": usage.get("guide_name") or "",
                "role_type": product.role_type.value if product else "",
                "form": usage.get("form") or first_nonempty(calc, "Type"),
                "category": (product.category if product else "") or first_nonempty(calc, "Cateogry", "Category"),
                "analysis": usage.get("analysis") or "",
                "npk": first_nonempty(calc, "NPK"),
                "product_url": (getattr(product, "product_url", None) or "")
                or first_nonempty(calc, "product_url", "Buy now"),
                "pack_sizes": first_nonempty(calc, "pack_sizes"),
                "calendar_how_often": interval.get("how_often") or "",
                "calendar_cadence_days": interval.get("cadence_days") or "",
                "calendar_window": interval.get("window") or "",
                "calendar_one_off": interval.get("one_off") or "0",
                "calendar_max_per_year": interval.get("max_per_year") or "",
                "calendar_tank_group": interval.get("tank_group") or "",
                "calendar_method": interval.get("method") or "",
                "calendar_notes": calendar_notes,
                "mix_together": pick_flag(interval, usage, "mix_together"),
                "independent": pick_flag(interval, usage, "independent"),
                "water_in": pick_flag(interval, usage, "water_in"),
                "soil_drench": pick_flag(interval, usage, "soil_drench"),
                "foliar": pick_flag(interval, usage, "foliar"),
                "fertigation": pick_flag(interval, usage, "fertigation"),
                "hydroponic": pick_flag(interval, usage, "hydroponic"),
                "hose_on": pick_flag(interval, usage, "hose_on"),
                "mix_with_iron": pick_flag(interval, usage, "mix_with_iron"),
                "usage_frequency_text": usage.get("frequency_text") or "",
                "usage_rate_text": usage.get("rate_text") or "",
                "usage_min_dilution": usage.get("min_dilution") or "",
                "usage_application_notes": usage_notes,
                "calculator_frequency": first_nonempty(calc, "calculator_frequency", "Frequency"),
                "calculator_dilution": first_nonempty(calc, "calculator_dilution", "Dilution Rate"),
                "calculator_rate_100m2": first_nonempty(calc, "rate_100m2"),
                "calculator_notes": calc_notes,
                "garden_natives_half_strength": natives,
                "combined_notes": combined,
                "in_catalog": "1" if product else "0",
                "in_calendar": "1" if interval else "0",
                "in_usage_guide": "1" if usage else "0",
                "in_rates_calculator": "1" if calc else "0",
            }
        )
    return products, guide_rates, calc_rows


def write_csv(path: Path, cols: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in cols})


def autosize(ws, max_width: int = 42) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = 0
        for cell in col:
            val = "" if cell.value is None else str(cell.value)
            longest = max(longest, min(len(val), max_width))
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), max_width)


def write_xlsx(
    path: Path,
    products: list[dict[str, str]],
    intervals: list[dict[str, str]],
    guide_rates: list[dict[str, str]],
    calc_rows: list[dict[str, str]],
) -> None:
    if Workbook is None:
        raise RuntimeError("openpyxl is required to write the Excel workbook")
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="14532D")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )
    flag_fills = {
        "mix_together": PatternFill("solid", fgColor="C6E6C3"),
        "independent": PatternFill("solid", fgColor="F8E3B0"),
        "water_in": PatternFill("solid", fgColor="C5D8F0"),
        "soil_drench": PatternFill("solid", fgColor="F5D0C0"),
        "foliar": PatternFill("solid", fgColor="C6E6C3"),
        "fertigation": PatternFill("solid", fgColor="C5CAE8"),
        "hydroponic": PatternFill("solid", fgColor="B2EBEB"),
        "hose_on": PatternFill("solid", fgColor="D4F0C0"),
        "mix_with_iron": PatternFill("solid", fgColor="F8E3B0"),
        "garden_natives_half_strength": PatternFill("solid", fgColor="F8E3B0"),
    }

    def fill_sheet(ws, cols: list[str], rows: list[dict[str, str]], table_name: str) -> None:
        ws.append(cols)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in rows:
            ws.append([row.get(c, "") for c in cols])
        for r in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2), min_col=1, max_col=len(cols)):
            for cell in r:
                cell.alignment = wrap
                cell.border = thin
                header = cols[cell.column - 1]
                if header in FLAG_COLS and str(cell.value) == "1" and header in flag_fills:
                    cell.fill = flag_fills[header]
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 32
        if rows:
            table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(cols))}{ws.max_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        autosize(ws)

    ws_read = wb.active
    ws_read.title = "Read_me"
    ws_read["A1"] = "Plant Doctor — master application rates, intervals, and notes"
    ws_read["A1"].font = Font(bold=True, size=14, color="14532D")
    howto = [
        "",
        "Join key: sku  (same IDs as Product Guide, Rates Calculator, and Application Calendar).",
        "Load CSV with: csv.DictReader(open(path, encoding='utf-8-sig'))",
        "Flag columns are 0/1: mix_together, independent, water_in, soil_drench, foliar, fertigation, hydroponic, hose_on, mix_with_iron, garden_natives_half_strength.",
        "Calendar badge flags on the products sheet come from schedule_intervals.csv (the calendar source of truth).",
        "Sheets:",
        "  products — one row per SKU, with calendar cadence, badges, usage-guide notes, and summarised calculator rates.",
        "  calendar_intervals — full calendar cadence table.",
        "  usage_guide_rates — exploded rate lines from the Product Usage Guide PDF.",
        "  calculator_rates — every plant-type / method row from the rates calculator sheet.",
        "  program_notes — plan-level notes (winter, natives half-rate, iron spacing, FFR/Activ8 alternate).",
        "  legend — flag meanings.",
        "Do not treat this as a substitute for the label. Adjust to growth, weather, and a jar test.",
    ]
    for i, line in enumerate(howto, start=2):
        ws_read[f"A{i}"] = line
        ws_read[f"A{i}"].alignment = Alignment(wrap_text=True)
    ws_read.column_dimensions["A"].width = 118

    ws_prod = wb.create_sheet("products")
    fill_sheet(ws_prod, PRODUCT_COLS, products, "MasterProducts")
    ws_prod.column_dimensions["AL"].width = 48

    interval_cols = [
        "sku",
        "cadence_days",
        "one_off",
        "window",
        "max_per_year",
        "tank_group",
        "method",
        "mix_together",
        "independent",
        "water_in",
        "soil_drench",
        "foliar",
        "fertigation",
        "hydroponic",
        "hose_on",
        "how_often",
        "notes",
    ]
    fill_sheet(wb.create_sheet("calendar_intervals"), interval_cols, intervals, "CalendarIntervals")

    guide_rate_cols = ["sku", "guide_name", "use_case", "method", "amount_text", "notes"]
    fill_sheet(wb.create_sheet("usage_guide_rates"), guide_rate_cols, guide_rates, "UsageGuideRates")

    calc_preferred = [
        "sku",
        "Product Name",
        "pack_sizes",
        "NPK",
        "Cateogry",
        "Type",
        "Uses",
        "Plant Type",
        "Specific Plants",
        "Dilution Rate",
        "rate_100m2",
        "Frequency",
        "product_url",
        "Notes",
        "Dilution Guide",
    ]
    calc_seen = set(calc_preferred)
    calc_extra = []
    for row in calc_rows:
        for key in row:
            if key not in calc_seen:
                calc_extra.append(key)
                calc_seen.add(key)
    fill_sheet(
        wb.create_sheet("calculator_rates"),
        calc_preferred + calc_extra,
        calc_rows,
        "CalculatorRates",
    )

    note_cols = ["note_id", "applies_when", "sku", "text"]
    fill_sheet(wb.create_sheet("program_notes"), note_cols, PROGRAM_NOTES, "ProgramNotes")

    ws_leg = wb.create_sheet("legend")
    ws_leg.append(["column", "value", "meaning"])
    for cell in ws_leg[1]:
        cell.font = header_font
        cell.fill = header_fill
    for row in [
        ("sku", "catalog id", "Join key across Product Guide, Rates Calculator, and Calendar."),
        ("mix_together", "1", "Can be mixed together as concentrates."),
        ("independent", "1", "Apply independently — do not tank-mix as a concentrate."),
        ("water_in", "1", "Water in after application."),
        ("garden_natives_half_strength", "1", "On garden plans, use this product at half strength on natives and sensitive plants."),
        ("calendar_window", "growing_season | year_round", "growing_season skips typical temperate winter; year_round eases to about half cadence in June–August."),
        ("tank_group", "mixable | no_iron_mix | iron | granular", "How the calendar groups same-day applications."),
        ("combined_notes", "text", "Calendar notes + usage-guide application notes + calculator notes, de-duplicated."),
    ]:
        ws_leg.append(list(row))
    autosize(ws_leg, 80)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    products, guide_rates, calc_rows = build_products()
    intervals = read_dicts(DATA / "schedule_intervals.csv")
    write_csv(DATA / "pd_master_products.csv", PRODUCT_COLS, products)
    write_csv(
        DATA / "pd_master_program_notes.csv",
        ["note_id", "applies_when", "sku", "text"],
        PROGRAM_NOTES,
    )
    write_xlsx(
        DATA / "pd_master_application_guide.xlsx",
        products,
        intervals,
        guide_rates,
        calc_rows,
    )
    print(f"products {len(products)}")
    print(f"usage-guide rate lines {len(guide_rates)}")
    print(f"calculator rate rows {len(calc_rows)}")
    print(DATA / "pd_master_products.csv")
    print(DATA / "pd_master_program_notes.csv")
    print(DATA / "pd_master_application_guide.xlsx")


if __name__ == "__main__":
    main()
